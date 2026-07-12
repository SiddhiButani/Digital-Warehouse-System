from flask import Blueprint, render_template, send_file, flash, redirect, url_for, session
import os
import csv
import io
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.units import inch

from routes.auth import login_required
from config import Config
from database.db import get_db_connection
from models.product import Product
from models.supplier import Supplier
from models.inventory import Inventory

reports_bp = Blueprint('reports', __name__)

@reports_bp.route('/')
@login_required()
def index():
    return render_template('reports.html', active_page='reports')

# Helper to fetch report data
def get_report_data(report_type):
    conn = get_db_connection()
    if report_type == 'inventory':
        headers = ['SKU ID', 'Product Name', 'Category', 'Stock Quantity', 'Unit Price ($)', 'Total Value ($)', 'Shelf Location', 'Supplier']
        query = """
            SELECT p.id, p.product_name, p.category, p.quantity, p.unit_price, 
                   (p.quantity * p.unit_price) AS total_val, p.shelf_id, s.company_name
            FROM products p
            LEFT JOIN suppliers s ON p.supplier_id = s.id
            ORDER BY p.id ASC;
        """
        rows = [list(r) for r in conn.execute(query).fetchall()]
        title = "Warehouse Inventory Balance Report"
    elif report_type == 'transactions':
        headers = ['Tx ID', 'SKU ID', 'Product Name', 'Action Type', 'Amount', 'Date & Time', 'Operator']
        query = """
            SELECT t.id, t.product_id, p.product_name, t.transaction_type, t.quantity, t.date, u.username
            FROM inventory_transactions t
            JOIN products p ON t.product_id = p.id
            JOIN users u ON t.user_id = u.id
            ORDER BY t.date DESC;
        """
        rows = [list(r) for r in conn.execute(query).fetchall()]
        title = "Inventory Transaction Audit Trail"
    elif report_type == 'suppliers':
        headers = ['Supplier ID', 'Company Name', 'Contact Person', 'Phone Number', 'Email Address', 'Warehouse Address']
        query = "SELECT id, company_name, contact_person, phone, email, address FROM suppliers ORDER BY company_name ASC;"
        rows = [list(r) for r in conn.execute(query).fetchall()]
        title = "Registered Supplier Partners Directory"
    else:
        headers, rows, title = [], [], ""
    
    conn.close()
    return headers, rows, title

@reports_bp.route('/export/<fmt>/<report_type>')
@login_required()
def export_report(fmt, report_type):
    if report_type not in ['inventory', 'transactions', 'suppliers']:
        flash("Invalid report selection.", "error")
        return redirect(url_for('reports.index'))
        
    headers, rows, title = get_report_data(report_type)
    if not headers:
        flash("Could not compile data for report.", "error")
        return redirect(url_for('reports.index'))

    filename = f"{report_type}_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}"

    # 1. EXPORT TO CSV
    if fmt == 'csv':
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(headers)
        writer.writerows(rows)
        output.seek(0)
        
        return send_file(
            io.BytesIO(output.getvalue().encode('utf-8')),
            mimetype='text/csv',
            as_attachment=True,
            download_name=f"{filename}.csv"
        )

    # 2. EXPORT TO EXCEL (OpenPyXL)
    elif fmt == 'excel':
        wb = Workbook()
        ws = wb.active
        ws.title = report_type.capitalize()
        
        # Enable grid lines explicitly
        ws.views.sheetView[0].showGridLines = True
        
        # Styles
        title_font = Font(name='Segoe UI', size=16, bold=True, color='1F2937')
        header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='4F46E5', end_color='4F46E5', fill_type='solid') # Indigo header
        data_font = Font(name='Segoe UI', size=10, color='374151')
        border_thin = Side(border_style="thin", color="E5E7EB")
        cell_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)
        
        # Add Title Row
        ws.append([title])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        ws.cell(row=1, column=1).font = title_font
        ws.cell(row=1, column=1).alignment = Alignment(vertical='center')
        ws.row_dimensions[1].height = 40
        
        # Empty space row
        ws.append([])
        
        # Add headers row
        ws.append(headers)
        ws.row_dimensions[3].height = 26
        
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = cell_border
            
        # Add Data rows
        for row_idx, row in enumerate(rows, 4):
            ws.append(row)
            ws.row_dimensions[row_idx].height = 20
            for col_idx, val in enumerate(row, 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = data_font
                cell.border = cell_border
                
                # Format numbers specifically
                if isinstance(val, (int, float)):
                    cell.alignment = Alignment(horizontal='right', vertical='center')
                    # If unit price or value, format as currency
                    if 'price' in headers[col_idx-1].lower() or 'value' in headers[col_idx-1].lower():
                        cell.number_format = '$#,##0.00'
                    else:
                        cell.number_format = '#,##0'
                else:
                    cell.alignment = Alignment(horizontal='left', vertical='center')
                    
        # Auto-adjust column widths
        for col in ws.columns:
            max_len = 0
            for cell in col:
                if cell.row < 3: # Skip title row length checks
                    continue
                val_str = str(cell.value or '')
                if len(val_str) > max_len:
                    max_len = len(val_str)
            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
            
        # Save workbook to memory buffer
        out_buf = io.BytesIO()
        wb.save(out_buf)
        out_buf.seek(0)
        
        return send_file(
            out_buf,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=f"{filename}.xlsx"
        )

    # 3. EXPORT TO PDF (ReportLab)
    elif fmt == 'pdf':
        out_buf = io.BytesIO()
        # Create SimpleDocTemplate document
        doc = SimpleDocTemplate(
            out_buf,
            pagesize=letter,
            rightMargin=0.5*inch, leftMargin=0.5*inch,
            topMargin=0.75*inch, bottomMargin=0.75*inch
        )
        
        # Styles setup
        styles = getSampleStyleSheet()
        
        title_style = ParagraphStyle(
            'DocTitle',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=20,
            leading=24,
            textColor=colors.HexColor('#1E3A8A'), # Navy Blue
            spaceAfter=15
        )
        
        meta_style = ParagraphStyle(
            'DocMeta',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=10,
            textColor=colors.HexColor('#4B5563'),
            spaceAfter=25
        )
        
        table_cell_style = ParagraphStyle(
            'CellText',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=8.5,
            leading=10,
            textColor=colors.HexColor('#374151')
        )
        
        table_header_style = ParagraphStyle(
            'HeaderText',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=9.5,
            textColor=colors.white
        )

        elements = []
        
        # Header block
        elements.append(Paragraph(title, title_style))
        meta_text = f"Report Compiled: {datetime.now().strftime('%B %d, %Y at %H:%M:%S')} | Operator: {session.get('username', 'System')}"
        elements.append(Paragraph(meta_text, meta_style))
        elements.append(Spacer(1, 10))
        
        # Convert header titles and cell strings into Paragraph components for auto text wrapping
        pdf_headers = [Paragraph(h, table_header_style) for h in headers]
        pdf_rows = []
        for row in rows:
            pdf_row = []
            for col_idx, col in enumerate(row):
                # Format floats or ints nicely
                if isinstance(col, float):
                    if 'price' in headers[col_idx].lower() or 'value' in headers[col_idx].lower():
                        val_str = f"${col:,.2f}"
                    else:
                        val_str = f"{col:,.2f}"
                elif isinstance(col, int):
                    val_str = f"{col:,}"
                else:
                    val_str = str(col or '')
                pdf_row.append(Paragraph(val_str, table_cell_style))
            pdf_rows.append(pdf_row)
            
        table_data = [pdf_headers] + pdf_rows
        
        # Determine column widths based on size
        col_count = len(headers)
        avail_width = 7.5 * inch # letter width 8.5" minus margins 1.0"
        col_width = avail_width / col_count
        
        # Table layout definitions
        t = Table(table_data, colWidths=[col_width]*col_count)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#4F46E5')), # Purple/Indigo Header
            ('ALIGN', (0,0), (-1,-1), 'LEFT'),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('TOPPADDING', (0,0), (-1,-1), 8),
            ('BOTTOMPADDING', (0,0), (-1,-1), 8),
            ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.HexColor('#F9FAFB'), colors.white]), # Striped rows
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E5E7EB')),
        ]))
        
        elements.append(t)
        
        # Build Document
        doc.build(elements)
        out_buf.seek(0)
        
        return send_file(
            out_buf,
            mimetype='application/pdf',
            as_attachment=True,
            download_name=f"{filename}.pdf"
        )
        
    flash("Report generation format not supported.", "error")
    return redirect(url_for('reports.index'))
